import copy, re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

d = docx.Document('template.docx')
body = d.element.body
SH = 'shots/'
COLAB = 'https://colab.research.google.com/drive/1b750JkYeiisxDnfrFAzUkW90f1I-YSvb'
DATE = '27 Oct 2026'

# ---------------- helpers ----------------
def set_lines(p, after=120, before=0, line=276):
    pf = p.paragraph_format
    pf.space_after = Pt(after / 20); pf.space_before = Pt(before / 20); pf.line_spacing = line / 240

def run(p, text, b=False, i=False, size=12):
    r = p.add_run(text); r.bold = b; r.italic = i; r.font.size = Pt(size)
    r.font.name = 'Times New Roman'; r.font.color.rgb = RGBColor(0, 0, 0)
    return r

def rich(p, text, size=12):
    # **bold** and *italic* markup
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part: continue
        if part.startswith('**'): run(p, part[2:-2], b=True, size=size)
        elif part.startswith('*'): run(p, part[1:-1], i=True, size=size)
        else: run(p, part, size=size)

def para(text, align=None, size=12, after=120):
    p = d.add_paragraph(); rich(p, text, size); set_lines(p, after)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def h1(text, page_break=False):
    p = d.add_paragraph(style='Heading 1')
    if page_break: p.paragraph_format.page_break_before = True
    run(p, text, b=True, size=14); set_lines(p, after=120, before=200, line=240)
    p.paragraph_format.keep_with_next = True
    return p

def bullets(items):
    for t in items:
        p = d.add_paragraph(style='List Paragraph'); rich(p, '•  ' + t); set_lines(p, after=60)
        p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for e in ('top', 'left', 'bottom', 'right'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '000000'); b.append(x)
    tcPr.append(b)
    m = OxmlElement('w:tcMar')
    for e, w in (('top', 50), ('left', 90), ('bottom', 50), ('right', 90)):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:w'), str(w)); x.set(qn('w:type'), 'dxa'); m.append(x)
    tcPr.append(m)

def caption(text):
    p = d.add_paragraph(); run(p, text, size=11); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_lines(p, after=200, before=60, line=240); return p

def table(title, header, rows, widths, size=10.5):
    cp = d.add_paragraph(); run(cp, title, size=11); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_lines(cp, after=60, before=120, line=240); cp.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1 + len(rows), cols=len(header)); t.autofit = False
    for r_i, vals in enumerate([header] + rows):
        for c_i, v in enumerate(vals):
            c = t.cell(r_i, c_i); c.width = Inches(widths[c_i]); cell_border(c)
            p = c.paragraphs[0]; run(p, str(v), b=(r_i == 0), size=size); set_lines(p, after=0, line=240)
            if c_i > 0 and r_i > 0 and re.fullmatch(r'[-−0-9.±()% ]+', str(v)): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); tr.append(h)
    for row in t.rows:
        rp = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); rp.append(cs)
    grid = t._tbl.tblGrid
    for gc, wd in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(wd * 1440)))
    tblW = t._tbl.tblPr.find(qn('w:tblW'))
    if tblW is None: tblW = OxmlElement('w:tblW'); t._tbl.tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(sum(widths) * 1440))); tblW.set(qn('w:type'), 'dxa')
    for row in (t.rows[:-1] if len(rows) <= 6 else []):
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True
    sp = d.add_paragraph(); set_lines(sp, after=80, line=240)
    return t

FIG = [0]
def fig(name, text, max_w=6.2, max_h=3.4, prefix=None):
    im = Image.open(SH + name + '.png'); ar = im.height / im.width
    w = min(max_w, max_h / ar)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_lines(p, after=0, line=240)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(SH + name + '.png', width=Inches(w))
    if prefix: label = prefix
    else: FIG[0] += 1; label = f'Figure {FIG[0]}'
    caption(f'{label}: {text}')
    return label

# ---------------- cover page ----------------
P = d.paragraphs
def replace_para(p, text, b=None, size=None):
    rs = p.runs
    for r in rs[1:]: r._r.getparent().remove(r._r)
    rs[0].text = text
    if b is not None: rs[0].bold = b
P[0].runs[0].text = 'Diabetes Progression Regression'
for r in P[0].runs[1:]: r._r.getparent().remove(r._r)

cover = d.tables[0]
vals = ['Rachel Elizabeth Joy', '2667213', 'rachel.elizabeth@mtech.christuniversity.in', None,
        '17 – Diabetes Progression Regression', 'scikit-learn: sklearn.datasets.load_diabetes (Efron et al., 2004)', COLAB, DATE]
for i, v in enumerate(vals):
    if v is None: continue
    c = cover.cell(i, 1); p = c.paragraphs[0]
    for r in p.runs[1:]: r._r.getparent().remove(r._r)
    if p.runs: p.runs[0].text = v; p.runs[0].italic = False
    else: run(p, v, size=11)

# signature line
sig = P[6]
for r in sig.runs: r.text = ''
sig.runs[0].text = 'Signature / typed name: Rachel Elizabeth Joy          Date: ' + DATE

ABSTRACT = ("This project predicts a quantitative measure of diabetes disease progression one year after baseline for 442 "
    "patients, using ten baseline measurements from the scikit-learn diabetes dataset. After exploratory analysis, the data were "
    "split 60/20/20 into training, validation and test sets, and all imputation, encoding, scaling and two engineered features "
    "were fitted inside a pipeline on training data only. A mean-predictor baseline was compared with ordinary least squares, "
    "Ridge, Lasso and ElasticNet (tuned with repeated 5-fold cross-validation), a random forest and gradient boosting. The "
    "regularised linear models performed best; Lasso (α = 0.398) was selected and reached a test RMSE of 52.2, MAE of 41.6 and "
    "R² of 0.548, against a baseline RMSE of 77.7. Coefficient paths show that Ridge shrinks all twelve coefficients smoothly, "
    "while Lasso drops features one at a time and keeps only the obesity–blood pressure interaction and serum triglycerides at "
    "strong regularisation.")
# abstract: P[8] guide (italic) -> remove ; P[9] placeholder -> abstract
P[8]._p.getparent().remove(P[8]._p)
ab = P[9]
for r in ab.runs[1:]: r._r.getparent().remove(r._r)
ab.runs[0].text = ABSTRACT; ab.runs[0].italic = False; ab.runs[0].font.size = Pt(12)
ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; set_lines(ab, after=120)

# ---------------- remove everything after the abstract (guide page and placeholders) ----------------
kids = list(body)
start = kids.index(ab._p) + 1
for el in kids[start:]:
    if el.tag != qn('w:sectPr'): body.remove(el)
# keep sectPr last: python-docx add_paragraph inserts before sectPr automatically

# =========================================================================================
h1('1. Problem statement and objective', page_break=True)
para("The task is a supervised regression problem. For each patient we predict a single number, a quantitative measure of "
     "diabetes disease progression recorded one year after a baseline visit, from ten measurements taken at that visit: age, sex, "
     "body mass index (BMI), average blood pressure and six blood serum values. The target is not a diagnosis; it is a "
     "continuous score that runs from 25 to 346, where higher values mean faster progression.")
para("A model like this could help a diabetes clinic or a care-planning team rank newly seen patients by expected risk, so that "
     "those likely to deteriorate fastest are offered earlier follow-up, dietary support or a medication review. It could also tell "
     "researchers which baseline measurements carry the most information about progression.")
para("The main metric is the **root mean squared error (RMSE)**. It is measured in the same units as the progression score, so "
     "it can be read directly as a typical size of error, and because it squares errors it penalises large misses more heavily, "
     "which matters when a badly under-predicted patient could miss timely care. MAE is reported alongside it as a more robust "
     "average error, and R² shows the share of variance explained compared with simply predicting the mean. Success means a "
     "clear improvement over the mean baseline on data the model has never seen.")

h1('2. Dataset description')
para("The data come from the study by Efron et al. (2004) and are shipped with scikit-learn as *load_diabetes* [1], [6]. "
     "Loading with *scaled=False* gives the original clinical units, so that scaling can be learnt inside the pipeline instead "
     "of being applied to the full dataset in advance. The serum columns s1 to s6 were renamed to their clinical meaning (Table 1).")
para("There are **442 rows and 11 columns**: ten features and the target. All columns are stored as float64 and there are **no "
     "missing values** (Figure 1). Sex is a binary category coded 1 and 2, with 235 and 207 patients respectively. The target has "
     "a mean of 152.1, a median of 140.5 and a standard deviation of 77.1, so a model that always predicts the mean would be off "
     "by about 77 points on average. The dataset is small, which strongly shapes the choices made later: complex models have "
     "little data to learn from, and every split decision changes results noticeably, so cross-validation is repeated.")
table('Table 1: Features used in the model',
      ['Feature', 'Type', 'Meaning', 'Missing values'],
      [['age', 'numeric', 'Age in years', '0'], ['sex', 'categorical', 'Sex, coded 1 / 2', '0'],
       ['bmi', 'numeric', 'Body mass index', '0'], ['bp', 'numeric', 'Average blood pressure', '0'],
       ['tc (s1)', 'numeric', 'Total serum cholesterol', '0'], ['ldl (s2)', 'numeric', 'Low-density lipoproteins', '0'],
       ['hdl (s3)', 'numeric', 'High-density lipoproteins', '0'], ['tch (s4)', 'numeric', 'Total cholesterol / HDL', '0'],
       ['ltg (s5)', 'numeric', 'Log of serum triglycerides', '0'], ['glu (s6)', 'numeric', 'Blood sugar level', '0'],
       ['target', 'numeric', 'Progression score after one year', '0']],
      [1.1, 1.1, 2.9, 1.1])
fig('s01_load', 'Loading the data in original units, data.head() and data.info(): 442 rows, 11 float columns, no nulls', max_h=4.0)
para("Figure 1 confirms the shape, the renamed columns and that every column is complete, so no row has to be dropped.")

h1('3. Exploratory data analysis')
fig('s02_target', 'Histogram and box plot of the target, with its skewness', max_h=2.3)
para("Figure 2 shows that progression is right-skewed (skewness 0.44): most patients score between 87 and 212, with a long tail "
     "towards 346 and no extreme outliers beyond the whiskers. The skew is mild enough that a transformation of the target was not needed.")
fig('s03_corr', 'Correlation matrix of all features and the target (greyscale heatmap)', max_h=3.6)
para("In Figure 3, the features most correlated with progression are **bmi (r = 0.59)**, **ltg (0.57)** and **bp (0.44)**, while "
     "**hdl** is negatively correlated (−0.40), which fits its role as \"good\" cholesterol. Several features are strongly "
     "correlated with each other: tc and ldl (r = 0.90), and hdl and tch (−0.74).")
fig('s04_feat_hist', 'Distributions of the nine numeric features in original units', max_h=3.2)
para("Figure 4 shows that the features sit on very different scales, from ltg around 3 to 6 to tc around 100 to 300. This is "
     "why scaling is essential before regularisation: Ridge and Lasso penalise coefficient size, so without scaling the "
     "penalty would fall unevenly. Most features are roughly bell-shaped; tch is discrete because it was rounded in the original study.")
fig('s05_scatter', 'Target against bmi and ltg with a fitted line, and target by sex', max_h=2.3)
para("Figure 5 explains the target best: progression rises steadily with both bmi and ltg, but the scatter around each line is "
     "wide, so no single measurement predicts it well and an R² well below 1 should be expected. The two sexes have similar "
     "medians, so sex alone carries little signal.")
fig('s06_collinear', 'Total cholesterol against LDL, showing near-collinearity (r = 0.90)', max_h=2.5)
para("Figure 6 shows the collinearity directly. When two inputs move together, ordinary least squares can give them large "
     "coefficients of opposite sign that cancel out. This is the behaviour the focus requirement examines, because Ridge, "
     "Lasso and ElasticNet each resolve it in a different way.")

h1('4. Data preparation and feature engineering')
para("The data were split **before** any fitting. A first split held out 20% as the test set, and the remainder was split again "
     "75/25 into training and validation sets, giving **264 / 89 / 89 rows** (Figure 8). Because the target is continuous, the "
     "splits were stratified on target quintiles, so each set covers the full range of progression; the target means of the "
     "three sets are 152.3, 150.8 and 152.9. The seed is fixed at 42 throughout.")
para("All other steps sit inside one scikit-learn Pipeline (Figure 7), so they are learnt on the training folds only and "
     "re-applied, unchanged, to validation and test data. This prevents leakage of test information into the scaler. Two features "
     "were engineered from clinical reasoning. **bmi_x_bp** multiplies BMI by blood pressure to represent combined metabolic load, "
     "since obesity and hypertension tend to worsen each other; it is correlated 0.63 with the target on the training set, more "
     "strongly than either input alone. **ldl_hdl_ratio** is a standard lipid-risk ratio (correlation 0.26). No outliers were removed: "
     "the box plots showed none that looked like recording errors, and with 442 patients each row is valuable.")
table('Table 2: Data preparation steps',
      ['Step', 'What I did', 'Why'],
      [['Missing values', 'Median imputer (numeric), most-frequent imputer (sex), inside the pipeline', 'No values are missing now, but the pipeline stays safe for new patients with gaps'],
       ['Encoding', 'One-hot encoding of sex with the first level dropped', 'Sex is a category, not a quantity; dropping one level avoids a redundant column'],
       ['Scaling', 'StandardScaler on all numeric and engineered features', 'Regularisation penalises coefficient size, so all features must be on one scale'],
       ['New features', 'bmi_x_bp, ldl_hdl_ratio', 'Combined metabolic load and lipid risk ratio, both clinically motivated'],
       ['Split', '264 train / 89 validation / 89 test, stratified on target quintiles, seed 42', 'Test set kept untouched; stratification keeps the target range similar']],
      [1.2, 2.5, 2.5])
fig('s07_prep_code', 'Feature engineering function and the preprocessing ColumnTransformer inside the pipeline', max_h=2.3)
fig('s08_split', 'Stratified 60/20/20 split and the resulting set sizes and target means', max_h=2.1)

h1('5. Models and experimental setup')
para("Seven models were compared, all wrapped in the same pipeline so that preprocessing was identical (Figure 9). The **baseline** "
     "is a DummyRegressor that always predicts the training mean. **Ordinary least squares (OLS)** is the unregularised linear "
     "reference. **Ridge** adds an L2 penalty that shrinks all coefficients towards zero [3]; **Lasso** adds an L1 penalty that can "
     "set coefficients exactly to zero [4]; **ElasticNet** mixes the two through l1_ratio [5]. Two tree ensembles, a **random "
     "forest** [7] and **gradient boosting** [8], test whether non-linear models add anything on so few rows.")
para("Tuning used **5-fold cross-validation repeated 3 times** (15 fits per setting) on the training set only, with negative RMSE "
     "as the score. Repetition was used because single 5-fold estimates vary by several points on 264 rows. Ridge was searched over "
     "61 values of α from 10⁻³ to 10³, Lasso over 51 values from 10⁻³ to 10², and ElasticNet over 26 values of α crossed with five "
     "l1_ratio values (Figure 9). The ensembles were tuned with a 25-draw randomised search over depth, number of trees, leaf "
     "size, feature fraction and, for boosting, learning rate and subsampling (Figure 10). The same RepeatedKFold object with "
     "seed 42 was used for every model, so all were compared on identical folds.")
para("The validation set played a separate role. It was never used for tuning, only as a second, independent check after cross-validation, "
     "so that a model that happened to fit the training folds well could be caught before the test set was opened. Model selection followed a "
     "rule fixed in advance: lowest cross-validated RMSE among the learnt models, with validation RMSE used only to confirm the ranking. "
     "This matters because the gaps between the best models are small, and choosing on the test set would make the reported score optimistic.")
table('Table 3: Models, chosen settings and cross-validated RMSE (training set)',
      ['Model', 'Key settings (after tuning)', 'CV RMSE', 'Notes'],
      [['Baseline: mean', 'Predicts training mean', '76.72 ± 5.57', 'Floor to beat'],
       ['OLS', 'No penalty', '55.69 ± 4.88', 'Unstable coefficients'],
       ['Ridge', 'α = 19.95', '55.45 ± 4.83', 'Keeps all features'],
       ['Lasso', 'α = 0.398', '55.32 ± 5.04', 'Lowest CV RMSE; 9 features'],
       ['ElasticNet', 'α = 0.398, l1_ratio = 0.9', '55.39 ± 4.95', 'Close to Lasso'],
       ['Random forest', '200 trees, leaf 5, max_features 0.33', '56.71 ± 4.91', 'No gain over linear'],
       ['Gradient boosting', '400 trees, lr 0.01, depth 2, subsample 0.7', '57.01 ± 5.93', 'Most variable']],
      [1.3, 2.6, 1.15, 1.25], size=10)
fig('s09_tune_linear', 'Repeated K-fold set-up and grid search for Ridge, Lasso and ElasticNet with the best settings', max_h=2.7)
fig('s10_tune_trees', 'Randomised search for the random forest and gradient boosting', max_h=2.3)

h1('6. Results')
para("Table 4 and Figure 11 compare every model on the same cross-validation folds and on the same validation set; the test "
     "set was not used for any of these choices. All learnt models clearly beat the baseline, cutting RMSE from about 77 to about 56. "
     "The three regularised linear models are best and lie within 0.1 of each other in cross-validated RMSE, well inside one "
     "standard deviation (about 5), so they are effectively tied. The tree ensembles are slightly worse, which suggests that the "
     "relationship is mostly linear once the bmi × bp interaction is supplied, and that 264 rows are too few for trees to find "
     "anything extra (Figure 12).")
table('Table 4: Model comparison on cross-validation and the validation set',
      ['Model', 'CV RMSE', 'Val RMSE', 'Val MAE', 'Val R²'],
      [['Baseline (mean)', '76.72', '77.75', '65.64', '0.000'], ['OLS', '55.69', '56.49', '46.28', '0.472'],
       ['Ridge', '55.45', '55.84', '45.31', '0.484'], ['Lasso', '55.32', '55.88', '45.44', '0.483'],
       ['ElasticNet', '55.39', '55.79', '45.40', '0.485'], ['Random forest', '56.71', '57.06', '46.57', '0.461'],
       ['Gradient boosting', '57.01', '56.91', '46.41', '0.464'],
       ['Lasso – final, TEST set', '—', '52.21', '41.63', '0.548'], ['Baseline – TEST set', '—', '77.65', '—', '—']],
      [2.1, 1.0, 1.0, 1.0, 1.0])
fig('s11_results', 'Evaluation loop over all models and the results table printed in the notebook', max_h=3.0)
fig('s12_compare', 'Cross-validated RMSE (mean ± sd) and validation RMSE for every model', max_h=2.6)
para("**Lasso** was chosen because it has the lowest cross-validated RMSE, the most stable estimate available, and because it "
     "gives a simpler model with three coefficients set to zero. It was refit on training plus validation data and scored **once** "
     "on the test set (Figure 13): **RMSE 52.21, MAE 41.63 and R² 0.548**, against a baseline RMSE of 77.65. The test RMSE is "
     "slightly lower than the CV estimate, which is within normal fold-to-fold variation for 89 patients rather than a sign of leakage.")
fig('s13_final_test', 'Selecting the model, refitting on train + validation and the single test-set evaluation', max_h=2.3)

h1('7. Focus requirement')
para("**Requirement:** *Compare Ridge, Lasso and ElasticNet and show how coefficients shrink as regularisation grows.*")
para("To isolate the effect of the penalty, the preprocessing pipeline was fitted on the training set once, and each model was refit "
     "on the same standardised matrix across 80 values of α from 10⁻³ to 10³. ElasticNet used the tuned l1_ratio of 0.9. The "
     "coefficients at each model's CV-chosen α were tabulated against OLS (Figure 14), the full paths were plotted (Figure 15), "
     "and the overall size and sparsity of the coefficient vector were tracked (Figure 16).")
fig('s14_focus_code', 'Fitting coefficient paths and the coefficients at each CV-chosen α compared with OLS', max_h=3.4)
para("The OLS column in Figure 14 shows the problem regularisation solves. Because bmi, bp and bmi_x_bp are strongly related, OLS "
     "gives bmi_x_bp a coefficient of 102 while giving bmi and bp large *negative* coefficients (about −40 each), which make no "
     "clinical sense. At their tuned α, all three penalised models remove these sign flips.")
fig('s15_paths', 'Coefficient paths of Ridge, Lasso and ElasticNet against α (log scale); dotted line marks the CV-chosen α', max_h=2.4)
para("Figure 15 shows the three different shrinkage patterns. **Ridge** shrinks every coefficient smoothly and proportionally; "
     "none reaches zero, so all twelve are kept even at α = 1000, and correlated features end up sharing weight (bmi 12.9, "
     "bmi_x_bp 19.5). **Lasso** shrinks in steps and drives features to exactly zero one at a time. At its chosen α of 0.398 it has "
     "already removed bmi, ldl and tch, because bmi_x_bp and tc already carry that information. **ElasticNet** behaves in between: "
     "with l1_ratio 0.9 its paths look like Lasso but are smoother, and at the chosen α it drops only tch and keeps bmi (9.9).")
fig('s16_norms', 'L2 norm of the coefficient vector and number of non-zero coefficients as α grows', max_h=2.2)
fig('s17_lasso_drop', 'Order in which Lasso removes features, and CV RMSE at selected α values', max_h=2.6)
para("Figure 16 summarises this. The coefficient norm falls for all three models, but the non-zero count drops step by step only "
     "for Lasso and ElasticNet; for Ridge it stays at twelve. Figure 17 gives the order of removal: tch goes first (α ≈ 0.38), then "
     "bp, ldl_hdl_ratio, age, tc, glu, sex, ldl and hdl. The last three survivors are bmi, ltg and finally **bmi_x_bp** (α ≈ 51), "
     "matching the strongest correlations in Section 3. One subtle effect is that bmi is zero at α = 0.398 but re-enters between "
     "about 0.6 and 18 as bmi_x_bp shrinks. Lasso paths need not be monotone when features are collinear. Finally, the CV error is "
     "almost flat (55.3 to 55.7) for α between 0.001 and 1, then rises to the baseline 76.7 at α = 100 (Appendix B). "
     "Moderate regularisation therefore improves interpretability and stability more than raw accuracy.")

h1('8. Error analysis and interpretation')
para("Figure 18 plots the single set of test predictions made in Section 6; nothing was tuned after this point. The points follow "
     "the diagonal, but the model **compresses predictions towards the mean**. Patients in the lowest quartile of actual progression "
     "are over-predicted by 44 points on average, and those in the highest quartile are under-predicted by 38.")
fig('s18_resid', 'Predicted versus actual and residuals versus predicted on the test set, with error by target quartile', max_h=3.2)
fig('s19_importance', 'Permutation importance on the validation set (model fitted on training data only)', max_h=2.3)
para("Permutation importance was computed on the **validation** set, to avoid a second look at the test set (Figure 19). Shuffling "
     "**ltg** or **bmi** raises RMSE by about 12 points each. Shuffling raw bmi also breaks bmi_x_bp, so bmi stays important even "
     "though Lasso set its own coefficient to zero. Blood pressure adds about 3 points, hdl about 1.5 and sex under 1; age, glu, tch "
     "and ldl add nothing measurable. In practice, body weight, triglycerides and blood pressure carry almost all the predictive signal.")
fig('s20_worst', 'The eight test patients with the largest absolute errors', max_h=2.0)
para("The worst cases (Figure 20) show where the model fails. Patient 380 progressed only to 52 but was predicted 176, mainly "
     "because of high blood pressure (111.7) and an above-median BMI; patient 256, with a BMI of 41.3 against a training median of 25.6, "
     "reached 346 but was predicted 244. These patients behave differently from what their baseline measurements suggest, "
     "probably because of factors the dataset does not record, such as medication, diet, adherence or HbA1c. More informative "
     "features, an interaction between ltg and bmi, or a model that predicts a range rather than one number (quantile regression) "
     "would address this better than further tuning. In short, the model is useful for ranking patients by broad risk band, "
     "but its individual predictions are too uncertain to set a precise expected score for one patient.")

h1('9. Ethics, limitations and risks')
para("The dataset contains 442 de-identified patients from a single study, with no record of ethnicity, location, income or "
     "treatment. The model may therefore not transfer to other populations, and its errors could fall unevenly on groups that are "
     "missing from the data. The largest harm comes from under-predicting high-risk patients: the model under-estimates the top "
     "quartile by about 38 points, so a patient who most needs early care could be ranked as average. Over-prediction for "
     "low-risk patients could cause unnecessary worry or treatment. Sex is used as a feature; it adds little signal, and its inclusion "
     "should be reviewed with clinicians.")
para("The model explains only about half of the variance, so it is suitable for research or as one input to a clinician's "
     "decision, never as an automatic triage rule. Limitations of the experiment itself include the small sample, a test set of only "
     "89 patients, the absence of any external validation, and the use of a single random split, so the reported scores carry "
     "uncertainty of several RMSE points. Any real deployment would need privacy safeguards, prospective validation and fairness "
     "checks on a larger, more diverse cohort.")

h1('10. Conclusion and future work')
para("Regularised linear regression worked best on this small, collinear dataset. Lasso reduced test RMSE from 77.7 for the baseline "
     "to 52.2 and explained 55% of the variance. The engineered bmi × bp feature became the strongest single coefficient. Tree "
     "ensembles did not help, and Ridge, Lasso and ElasticNet reached almost identical accuracy. Their real difference is how they "
     "shrink: Ridge spreads weight across correlated features, while Lasso selects a few. What did not work was predicting the "
     "extremes, where errors of around 100 points remain. With more time I would use nested cross-validation and bootstrap "
     "confidence intervals for the test score, try quantile regression or a log-transformed target for the tails, and test the model "
     "on an external cohort with richer clinical data.")

h1('References')
refs = [
 "B. Efron, T. Hastie, I. Johnstone and R. Tibshirani, \"Least angle regression,\" *Annals of Statistics*, vol. 32, no. 2, pp. 407–499, 2004.",
 "F. Pedregosa et al., \"Scikit-learn: Machine learning in Python,\" *Journal of Machine Learning Research*, vol. 12, pp. 2825–2830, 2011.",
 "A. E. Hoerl and R. W. Kennard, \"Ridge regression: Biased estimation for nonorthogonal problems,\" *Technometrics*, vol. 12, no. 1, pp. 55–67, 1970.",
 "R. Tibshirani, \"Regression shrinkage and selection via the lasso,\" *Journal of the Royal Statistical Society, Series B*, vol. 58, no. 1, pp. 267–288, 1996.",
 "H. Zou and T. Hastie, \"Regularization and variable selection via the elastic net,\" *Journal of the Royal Statistical Society, Series B*, vol. 67, no. 2, pp. 301–320, 2005.",
 "scikit-learn developers, \"sklearn.datasets.load_diabetes,\" scikit-learn documentation. [Online]. Available: https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_diabetes.html",
 "L. Breiman, \"Random forests,\" *Machine Learning*, vol. 45, pp. 5–32, 2001.",
 "J. H. Friedman, \"Greedy function approximation: A gradient boosting machine,\" *Annals of Statistics*, vol. 29, no. 5, pp. 1189–1232, 2001.",
]
for i, r_ in enumerate(refs, 1):
    p = para(f'[{i}] ' + r_, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=80)
    p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)

h1('Appendix A: Key code snippets', page_break=True)
fig('a1_setup', 'Imports, fixed seed and greyscale plot style', max_h=3.3, prefix='Figure A1')
fig('a2_feature_check', 'Check that scaling is fitted on training data (mean 0, sd 1) and engineered-feature correlations', max_h=3.3, prefix='Figure A2')
fig('a3_paths_code', 'Plotting the coefficient paths with greyscale line styles', max_h=3.3, prefix='Figure A3')
fig('a4_drop_code', 'Finding the α beyond which each Lasso coefficient stays at zero', max_h=3.3, prefix='Figure A4')
fig('a5_resid_code', 'Residual plots and error by target quartile', max_h=3.3, prefix='Figure A5')
fig('a6_importance_code', 'Permutation importance on the validation set with a model fitted on training data only', max_h=3.3, prefix='Figure A6')

h1('Appendix B: Notebook outputs', page_break=True)
fig('b1_runall', 'Notebook title cell and the final cell after Runtime → Run all, completed without errors', max_h=2.6, prefix='Figure B1')
fig('b2_describe', 'Summary statistics of every column in original units', max_h=2.4, prefix='Figure B2')
fig('b3_cvcurve', 'Cross-validated RMSE against α for Ridge, Lasso and ElasticNet', max_h=2.6, prefix='Figure B3')

h1('Appendix C: AI tool declaration')
table('Table C1: AI tools used',
      ['Tool', 'What I used it for', 'What I changed or checked myself'],
      [['Claude (Anthropic)', 'Drafting the notebook code structure, the plots and a first draft of the report text', '[Describe what you checked: e.g. re-ran the notebook in Colab, verified every number in the report against the outputs, rewrote sections in my own words]'],
       ['', '', '']],
      [1.4, 2.4, 2.4])

d.save('2667213_Rachel_Elizabeth_Joy_Project.docx')
print('saved')
